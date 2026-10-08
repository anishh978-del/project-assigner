"""Write the assignments out as a formatted .xlsx."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="16191F")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
ALT_FILL = PatternFill("solid", fgColor="F4F2EF")
THIN = Side(style="thin", color="DDD8D2")
BORDER = Border(bottom=THIN)


def _head(ws, headers):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _autosize(ws, maximum=60):
    for col in range(1, ws.max_column + 1):
        width = max((len(str(ws.cell(row=r, column=col).value or ""))
                     for r in range(1, ws.max_row + 1)), default=10)
        ws.column_dimensions[get_column_letter(col)].width = min(width + 3, maximum)


def write_report(path, assignments, summary, source_note=""):
    wb = Workbook()

    # --- one row per student ---
    ws = wb.active
    ws.title = "Assignments"
    _head(ws, ["Group", "Student Name", "USN", "Project #", "Project",
               "Set", "Language", "Services"])
    for a in assignments:
        p = a["project"]
        for m in a["members"]:
            ws.append([a["group"], m["name"], m["id"], p["id"], p["title"],
                       p["set"], p["language"], ", ".join(p["services"])])
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value % 2 == 0:
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).fill = ALT_FILL
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).border = BORDER
    _autosize(ws)

    # --- one row per group ---
    ws2 = wb.create_sheet("Groups")
    _head(ws2, ["Group", "Size", "Members", "Project #", "Project",
                "Language", "Starter provided", "The hard part"])
    for a in assignments:
        p = a["project"]
        ws2.append([a["group"], len(a["members"]),
                    ", ".join(m["name"] for m in a["members"]),
                    p["id"], p["title"], p["language"],
                    "yes" if p.get("starter") else "no", p["hard_part"]])
    for r in range(2, ws2.max_row + 1):
        ws2.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="top")
    _autosize(ws2, maximum=50)
    ws2.column_dimensions["C"].width = 42
    ws2.column_dimensions["H"].width = 50

    # --- what was run ---
    ws3 = wb.create_sheet("Summary")
    _head(ws3, ["Field", "Value"])
    for k, v in summary.items():
        ws3.append([k.replace("_", " ").title(), v])
    if source_note:
        ws3.append(["Source", source_note])
    _autosize(ws3)

    # create the folder if the caller asked for one that does not exist yet
    parent = Path(path).parent
    if str(parent) not in ("", "."):
        parent.mkdir(parents=True, exist_ok=True)

    wb.save(path)
    return path
